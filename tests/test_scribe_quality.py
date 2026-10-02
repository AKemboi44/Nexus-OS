from app.agents.scribe_agent import ScribeResearchAgent
import json
from types import SimpleNamespace
from docx import Document
from docx.shared import Inches
from app.synthesis.citation_engine import CitationEngine


def test_scribe_deduplicates_section_items():
    agent = ScribeResearchAgent()
    agent.client = None

    items = agent._unique_items([
        "Evidence supports careful evaluation.",
        "Evidence supports careful evaluation.",
        "Different wording should remain.",
    ])

    assert items == [
        "Evidence supports careful evaluation.",
        "Different wording should remain.",
    ]


def test_scribe_does_not_append_repeated_citation_or_template_prose():
    agent = ScribeResearchAgent()
    agent.client = None
    sources = [
        {
            "title": "Research quality study",
            "authors": ["Abdar"],
            "year": 2021,
            "abstract": "This study provides enough evidence detail for a grounded paragraph about research quality and implementation.",
        }
    ]

    paragraph = agent._evidence_paragraph(
        "Evidence supports careful evaluation. (Abdar, 2021)",
        sources,
    )

    assert paragraph.count("Abdar, 2021") == 1
    assert "The selected studies report related findings" not in paragraph
    assert "This distinction matters" not in paragraph


def test_scribe_paragraph_is_concise_and_evidence_bounded():
    agent = ScribeResearchAgent()
    agent.client = None
    sources = [
        {
            "title": "Research quality study",
            "authors": ["Abdar"],
            "year": 2021,
            "abstract": "This study provides enough evidence detail for a grounded paragraph about research quality and implementation.",
        }
    ]

    paragraph = agent._evidence_paragraph(
        "Evidence supports careful evaluation of the research problem.",
        sources,
        topic="research quality",
        section="Key Themes",
        report_type="proposal",
    )
    words = paragraph.split()
    citation_index = paragraph.index("(Abdar, 2021)")
    before = paragraph[:citation_index]
    after = paragraph[citation_index + len("(Abdar, 2021)"):]

    assert len(words) < 60
    assert before.strip().endswith("Evidence supports careful evaluation of the research problem")
    assert after == "."
    assert paragraph.count("(Abdar, 2021)") == 1
    assert "(Abdar, 2021)." in paragraph
    assert ". (Abdar, 2021)" not in paragraph


def test_scribe_keeps_long_claims_focused_without_stock_padding():
    agent = ScribeResearchAgent()
    agent.client = None
    sources = [
        {
            "title": "A very long source title describing an extensive multi-dimensional investigation of research quality, implementation, measurement, and transferability",
            "authors": ["Abdar"],
            "year": 2021,
            "abstract": "This study provides enough evidence detail for a grounded paragraph about research quality and implementation.",
        }
    ]

    paragraph = agent._evidence_paragraph(
        "This is a deliberately long synthesized theme that contains many methodological, contextual, operational, and analytical dimensions which must remain focused on the same research question.",
        sources,
        topic="research quality",
        section="Key Themes",
        report_type="proposal",
    )

    assert len(paragraph.split()) < 80
    assert paragraph.count(".") <= 3


def test_scribe_replaces_garbled_repetitive_abstract_with_clean_fallback():
    agent = ScribeResearchAgent()
    abstract = agent._format_abstract(
        "There is is is is a a a failure failure mode in in large large models.",
        "language model reliability",
        [],
    )

    assert "is is" not in abstract
    assert "a a" not in abstract
    assert "failure failure" not in abstract
    assert "thatwe" not in abstract
    assert "Keywords:" in abstract
    assert len(abstract.split()) >= 100


def test_scribe_prose_cleanup_handles_long_repeated_input():
    repeated_dash = "claim" + ("--" * 10000) + "supported"
    repeated_words = ("evidence " * 10000).strip()

    cleaned = ScribeResearchAgent._clean_prose(
        f"{repeated_dash}; {repeated_words}"
    )

    assert "--" not in cleaned
    assert "evidence evidence" not in cleaned
    assert cleaned.endswith("evidence")


def test_scribe_keeps_proposal_body_flush_left_and_paragraph_spacing_consistent():
    document = Document()
    ScribeResearchAgent._add_body_paragraph(
        document,
        "A developed research paragraph.\n\nA second developed paragraph.",
    )
    ScribeResearchAgent._add_body_paragraph(
        document,
        "An abstract paragraph.",
        first_line_indent=0,
    )

    assert len(document.paragraphs) == 3
    assert document.paragraphs[0].paragraph_format.first_line_indent == 0
    assert document.paragraphs[1].paragraph_format.first_line_indent == 0
    assert document.paragraphs[2].paragraph_format.first_line_indent == 0
    assert all(paragraph.paragraph_format.space_before.pt == 0 for paragraph in document.paragraphs)
    assert all(paragraph.paragraph_format.space_after.pt == 0 for paragraph in document.paragraphs)


def test_apa_references_format_author_initials_and_in_text_et_al():
    source = {
        "authors": ["Ahmed Abdelghany", "Khaled Abdelghany", "Laila Hassan"],
        "year": 2026,
        "title": "Research on evidence-grounded systems",
        "venue": "Journal of Applied Research",
        "doi": "10.1234/example",
    }
    agent = ScribeResearchAgent()

    reference = CitationEngine.generate_apa_7th(source)
    citation = agent._citation_for_sources([source])

    assert reference.startswith("Abdelghany, A., Abdelghany, K., & Hassan, L. (2026).")
    assert "Journal of Applied Research." in reference
    assert "https://doi.org/10.1234/example" in reference
    assert citation == "(Abdelghany et al., 2026)"


def test_apa_reference_document_italics_title_or_publication():
    document = Document()
    paragraph = document.add_paragraph()
    source = {
        "authors": ["MANSOUR, Amina"],
        "year": 2026,
        "title": "A study of responsible systems",
        "venue": "Journal of Research",
        "volume": "12",
        "issue": "2",
        "pages": "34-49",
    }

    ScribeResearchAgent._add_apa_reference(paragraph, source)

    assert paragraph.text == (
        "Mansour, A. (2026). A study of responsible systems. "
        "Journal of Research, 12(2), 34-49."
    )
    assert next(run for run in paragraph.runs if run.text == "Journal of Research").italic


def test_apa_reference_sorting_is_alphabetical_by_first_author():
    sources = [
        {"authors": ["Zoe Smith"], "year": 2024, "title": "Later"},
        {"authors": ["Ahmed Abdelghany"], "year": 2022, "title": "Earlier"},
    ]

    ordered = sorted(sources, key=ScribeResearchAgent._apa_reference_sort_key)

    assert [source["title"] for source in ordered] == ["Earlier", "Later"]


def test_proposal_editor_generates_full_grounded_structure_and_citations():
    agent = ScribeResearchAgent()
    source = {
        "title": "Research quality across settings",
        "authors": ["Atul Kumar"],
        "year": 2024,
        "venue": "Journal of Applied Research",
        "url": "https://example.org/research",
        "abstract": "The study compares research-quality measures across institutional settings and reports differences in implementation and evaluation.",
    }

    def paragraph(text, evidence_ids=None):
        return {
            "text": text,
            "evidence_ids": ["S1"] if evidence_ids is None else evidence_ids,
        }

    developed = (
        "The reviewed source compares research-quality measures across institutional settings and identifies "
        "differences in how teams implement and evaluate those measures. Its reported comparison provides "
        "a relevant basis for examining the topic, but the available abstract does not establish why the "
        "observed differences arose or whether they generalize beyond the settings studied. The proposal "
        "should therefore distinguish the documented comparison from explanations that require further "
        "empirical testing. A full-text review is needed to verify the measures, context, and limits of the finding."
    )
    methodology = (
        "A comparative design is a provisional option because the research problem concerns variation across "
        "settings rather than a single reported outcome. The study should specify the unit of analysis, the "
        "comparison groups, and the indicators used to define research quality before data collection begins. "
        "The supplied records do not identify a target population or a validated measurement instrument. "
        "These choices should remain explicit design decisions, not be presented as facts established by the "
        "reviewed study. [Specify the study population and setting.]"
    )

    class FakeModels:
        def generate_content(self, model, contents):
            assert "expert academic research consultant and scholar" in contents
            assert "zero" not in contents.lower() or "FACTUAL INTEGRITY" in contents
            assert "research-quality measures across institutional settings" in contents
            lit = {
                key: [paragraph(developed), paragraph(developed)]
                for key in (
                    "key_themes",
                    "contradictions_and_boundary_conditions",
                    "research_gaps",
                    "research_areas",
                    "opportunities",
                    "research_problems",
                )
            }
            methods = {
                key: [paragraph(methodology), paragraph(methodology, [])]
                for key in ("research_design", "data_collection", "analysis", "limitations")
            }
            response = {
                "introduction": [paragraph(developed), paragraph(developed)],
                "problem_statement": [paragraph(developed), paragraph(developed)],
                "research_questions": [
                    paragraph("How do research-quality measures vary across institutional settings?", ["S1"]),
                    paragraph("Which contextual conditions explain variation in the implementation of those measures?", ["S1"]),
                ],
                "hypotheses": [],
                "research_objectives": [
                    paragraph("Compare quality measures across the institutional settings represented in the evidence.", ["S1"]),
                    paragraph("Identify contextual conditions associated with differences in implementation.", ["S1"]),
                ],
                "literature_review": lit,
                "conceptual_framework": [
                    paragraph(developed),
                    paragraph(
                        "No established framework is identified in the supplied records. "
                        "The proposal should not attach an unrelated theory to the topic or "
                        "present an unverified model as established scholarship. Before selecting "
                        "a framework, the researcher should review peer-reviewed theoretical work "
                        "that directly addresses research quality across institutional settings. "
                        "The chosen framework should clarify rather than predetermine the proposed analysis. "
                        "This decision requires a focused search beyond the abstracts supplied here. "
                        "[Add a verified peer-reviewed source for the selected framework.]",
                        [],
                    ),
                ],
                "methodology": methods,
                "significance": [paragraph(developed), paragraph(developed)],
                "timeline": [
                    {"phase": "Protocol", "duration": "Months 1-2", "activities": "Confirm questions, scope, and eligibility criteria."},
                    {"phase": "Review", "duration": "Months 3-4", "activities": "Verify full texts and finalize evidence extraction."},
                    {"phase": "Collection", "duration": "Months 5-7", "activities": "Collect approved data using documented procedures."},
                    {"phase": "Analysis", "duration": "Months 8-10", "activities": "Analyze results, document limitations, and prepare the report."},
                ],
            }
            return SimpleNamespace(text=json.dumps(response))

    agent.client = SimpleNamespace(models=FakeModels())
    agent._prepare_proposal_draft(
        "research quality",
        [source],
        ["Measures vary across settings."],
        ["Implementation differs by institution."],
        ["The evidence does not identify causes."],
        ["Compare institutions."],
        ["Validate a shared measure."],
        ["Determine which measures are useful."],
        "scholarly",
    )
    assert agent._proposal_draft is not None
    assert agent._editorial_synthesis_complete
    assert agent._validate_proposal_draft(agent._proposal_draft, {}) is None

    document = Document()
    agent._write_proposal_sections(
        document,
        "research quality",
        [source],
        ["Measures vary across settings."],
        ["Implementation differs by institution."],
        ["The evidence does not identify causes."],
        ["Validate a shared measure."],
        ["Determine which measures are useful."],
    )
    paragraphs = document.paragraphs
    text = "\n".join(paragraph.text for paragraph in paragraphs)
    for heading in (
        "Introduction and Background",
        "Statement of the Problem",
        "Research Questions and Hypotheses",
        "Research Objectives",
        "Literature Review",
        "Key Themes",
        "Contradictions and Boundary Conditions",
        "Research Gaps",
        "Research Areas",
        "Opportunities",
        "Research Problems",
        "Theoretical or Conceptual Framework",
        "Research Design",
        "Data Collection",
        "Analysis Strategy",
        "Significance and Implications",
        "Preliminary Timeline",
    ):
        assert heading in text
    assert "(Kumar, 2024)" in text
    assert all(paragraph.paragraph_format.first_line_indent == 0 for paragraph in paragraphs)
    assert sum(paragraph.text.startswith("Phase ") for paragraph in paragraphs) == 4


def test_scribe_quotes_direct_topic_references_in_report_body_only():
    document = Document()
    document.add_paragraph("Research Proposal: Research quality")
    document.add_paragraph("Abstract")
    paragraph = document.add_paragraph()
    paragraph.add_run("This report examines research quality. ")
    paragraph.add_run("Research quality, specifically, remains important.")
    quoted = document.add_paragraph("The topic “research quality” is already quoted.")
    document.add_paragraph("References")
    reference = document.add_paragraph("Author (2024). Research quality study.")

    ScribeResearchAgent._quote_research_topic_references(document, "research quality")

    assert paragraph.text == (
        "This report examines “research quality”. "
        "“Research quality”, specifically, remains important."
    )
    assert quoted.text == "The topic “research quality” is already quoted."
    assert reference.text == "Author (2024). Research quality study."
    assert document.paragraphs[0].text == "Research Proposal: Research quality"


def test_scribe_formats_publication_style_abstract_with_keywords():
    agent = ScribeResearchAgent()
    agent.client = None
    abstract = agent._format_abstract(
        "This review examines token optimization and compares prompt compression with dynamic pruning.",
        "token optimization",
        [],
    )

    assert "Keywords:" in abstract
    assert len(abstract.split()) >= 80
    assert "This review examines" in abstract


def test_scribe_removes_parenthetical_year_citations_without_regex_backtracking():
    cleaned = ScribeResearchAgent._remove_parenthetical_years(
        "Evidence supports this claim (Author, 2021), but keeps this note (important context)."
    )

    assert cleaned == "Evidence supports this claim, but keeps this note (important context)."


def test_scribe_parses_keyword_label_with_optional_spacing():
    agent = ScribeResearchAgent()
    agent.client = None

    abstract = agent._format_abstract(
        "A sufficiently detailed review abstract discusses evidence across methods and settings. "
        "Keywords : research quality, validation",
        "research quality",
        [],
    )

    assert "Keywords: research quality, validation." in abstract


def test_synthesis_parser_reads_research_areas_and_inline_abstract_heading():
    from app.reports.dossier_generator import DossierGenerator
    from models.research_dossier import ResearchDossier

    dossier = ResearchDossier("research quality")
    DossierGenerator()._parse_synthesis_payload(
        "### ABSTRACT: This review evaluates research quality across study settings.\n"
        "### RESEARCH AREAS\n"
        "- Compare measurement approaches across the included study populations.\n"
        "### OPPORTUNITY AREAS\n"
        "- Test a practical quality-assessment workflow for research teams.\n",
        dossier,
    )

    assert dossier.abstract == "This review evaluates research quality across study settings."
    assert dossier.research_areas == [
        "Compare measurement approaches across the included study populations."
    ]
    assert dossier.opportunity_areas == [
        "Test a practical quality-assessment workflow for research teams."
    ]


def test_scribe_cleans_space_before_punctuation():
    cleaned = ScribeResearchAgent._remove_incomplete_fragments(
        "Evidence remains useful , despite limits ."
    )

    assert cleaned == "Evidence remains useful, despite limits."


def test_complete_literature_review_contains_paper_components():
    agent = ScribeResearchAgent()
    agent.client = None
    sections = []

    class FakeDoc:
        def add_paragraph(self):
            class Paragraph:
                alignment = None
                paragraph_format = type("Format", (), {})()

                def add_run(self, text=""):
                    sections.append(str(text))
                    return type(
                        "Run",
                        (),
                        {"font": type("Font", (), {})()},
                    )()

            return Paragraph()

    sources = [{
        "title": "Research quality study",
        "authors": ["Abdar"],
        "year": 2021,
        "abstract": "This study provides enough evidence detail for a grounded paragraph about research quality and implementation.",
    }]
    agent._write_full_starter_sections(
        FakeDoc(),
        "research quality",
        sources,
        ["A recurring evidence theme."],
        ["A contextual disagreement."],
        ["A methodological gap."],
        ["A practical opportunity."],
        ["A research problem."],
        "scholarly",
        research_areas=["Compare methods across the included studies."],
    )

    output = " ".join(sections)
    for heading in (
        "Research Objectives",
        "Recommended Research Areas",
        "Conceptual Framework",
        "Methodology",
        "Data Collection",
        "Synthesis Findings",
        "Discussion",
        "Proposed Research Direction",
        "Operationalization and Study Design",
        "Conclusion",
    ):
        assert heading in output


def test_report_editorial_synthesis_expands_each_theme_from_source_evidence():
    agent = ScribeResearchAgent()
    source = {
        "title": "Research quality study",
        "authors": ["Abdar"],
        "year": 2021,
        "abstract": "The study examines research quality and implementation across multiple settings.",
    }

    class FakeModels:
        calls = 0

        def generate_content(self, model, contents):
            self.calls += 1
            assert "exactly two" in contents
            payload = json.loads(contents.split("Report sections:", 1)[1])
            paragraphs = []
            for entry in payload:
                claim = entry["claim"].rstrip(".")
                first_paragraph = (
                    f"For the research topic, {claim} provides a focused way to interpret the evidence. "
                    "The supplied source records place this idea within a defined research context and "
                    "help distinguish it from an assumption that would apply to every setting. Read "
                    "together, the evidence supports considering how the relevant concepts relate. "
                    "However, the available abstracts do not establish that the same relationship holds "
                    "across all populations or methods. This qualification keeps the synthesis "
                    "proportionate to the material reviewed."
                )
                second_paragraph = (
                    "This pattern matters because differences in setting and measurement can change how "
                    "a research concept is understood and evaluated. The current source descriptions "
                    "provide a basis for comparison, but they do not resolve whether the interpretation "
                    "transfers beyond the studies represented here. Further work should identify the "
                    "population, measures, and contextual conditions needed to test the proposition. "
                    "Such a design would clarify its relevance without treating a promising interpretation "
                    "as an established result or assuming that it applies universally."
                )
                paragraphs.append({
                    "id": entry["id"],
                    "paragraphs": [first_paragraph, second_paragraph],
                })
            return SimpleNamespace(text=json.dumps({"paragraphs": paragraphs}))

    class FakeClient:
        models = FakeModels()

    agent.client = FakeClient()
    agent._prepare_editorial_synthesis(
        "research quality",
        [source],
        ["Quality measurement varies across research settings."],
        [],
        [],
        [],
        [],
        [],
        "proposal",
        "scholarly",
    )

    paragraph = agent._evidence_paragraph(
        "Quality measurement varies across research settings.",
        [source],
        topic="research quality",
        section="Key Themes",
    )
    assert agent.client.models.calls == 1
    assert agent._editorial_synthesis_complete is True
    assert len(paragraph.split()) >= 75
    assert paragraph.count("(Abdar, 2021)") == 2
    assert "Further work should identify the population" in paragraph
