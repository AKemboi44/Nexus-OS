"""Adaptive matching: when no source meets every selected criterion, the closest matches fill the gap and every
source says what it met and missed. A strict switch restores all-or-nothing."""

import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook

import cloud_app
from app.payments import exposure
from app.research import criteria
from app.research.criteria import (
    LABEL_CITED, LABEL_OPEN, LABEL_PEER, LABEL_RECENT, LABEL_REVIEW, LABEL_UNIQUE, evaluate, needed_for_closest,
    qualification, resolve,
)
from app.research.research_pipeline import ResearchPipeline
from models.research_dossier import ResearchDossier
from tests.test_free_tier_exposure import World, sources
from tests.test_criteria import NOW, src

@pytest.fixture(autouse=True)
def half_share(monkeypatch):
    monkeypatch.setenv(criteria.MIN_MATCH_ENV, "50")


SIX = [LABEL_PEER, LABEL_CITED, LABEL_UNIQUE, LABEL_RECENT, LABEL_OPEN, LABEL_REVIEW]


# --- evaluating one source ----------------------------------------------------------------------------------------------------------

def test_a_source_reports_what_it_met_and_what_it_missed_with_the_fact():
    result = evaluate(src(is_open_access=False), resolve([LABEL_PEER, LABEL_OPEN, LABEL_CITED], None, True), NOW)
    assert result.met == (LABEL_PEER, LABEL_CITED)
    assert result.missed == (f"{LABEL_OPEN} (not open access)",)
    assert (result.total, result.result_text, result.hard_missed) == (3, "Met 2 of 3", False)


def test_meeting_everything_reads_met_all():
    assert evaluate(src(), resolve([LABEL_PEER, LABEL_CITED], None, True), NOW).result_text == "Met all 2"


@pytest.mark.parametrize("total, needed", [(1, 1), (2, 1), (3, 2), (4, 2), (5, 3), (6, 3), (8, 4)])
def test_a_closest_match_must_meet_at_least_half_of_the_criteria(total, needed):
    assert needed_for_closest(total) == needed


def test_qualification_full_closest_and_no():
    resolved = resolve([LABEL_PEER, LABEL_CITED, LABEL_OPEN, LABEL_RECENT], None, True)
    assert qualification(evaluate(src(), resolved, NOW), resolved) == "full"
    assert qualification(evaluate(src(is_open_access=False), resolved, NOW), resolved) == "closest"
    assert qualification(evaluate(src(is_open_access=False, year=2010), resolved, NOW), resolved) == "closest"   # 2 of 4
    assert qualification(evaluate(src(is_open_access=False, year=2010, citation_count=0), resolved, NOW), resolved) == "no"


def test_a_single_criterion_has_no_closest_match():
    resolved = resolve([LABEL_OPEN], None, True)
    assert qualification(evaluate(src(is_open_access=False), resolved, NOW), resolved) == "no"


def test_the_publication_type_is_never_relaxed():
    resolved = resolve([LABEL_CITED, LABEL_RECENT, LABEL_OPEN], "journal", True)
    preprint = evaluate(src(work_type="preprint"), resolved, NOW)
    assert preprint.hard_missed is True and len(preprint.met) == 3
    assert qualification(preprint, resolved) == "no", "meets 3 of 4 but misses the required type"


def test_strict_matching_accepts_only_full_matches():
    strict = resolve([LABEL_PEER, LABEL_CITED, LABEL_OPEN], None, True, strict=True)
    assert strict.strict is True
    assert qualification(evaluate(src(is_open_access=False), strict, NOW), strict) == "no"
    assert qualification(evaluate(src(), strict, NOW), strict) == "full"


# --- inside the scan: the owner's scenario ------------------------------------------------------------------------------------------------

def pool():
    """Good, recent, cited, open journal articles, but none is a review: no source can meet all six criteria."""
    return [
        src(uid="a", title="Reasoning evaluation in language models", doi="10.1/a", citation_count=40),
        src(uid="b", title="Reasoning evaluation benchmarks", doi="10.1/b", citation_count=30, year=2025),
        src(uid="c", title="Reasoning evaluation closed models", doi="10.1/c", citation_count=12, is_open_access=False),
        src(uid="far", title="Reasoning evaluation old note", doi="10.1/far", year=2005, citation_count=0, is_open_access=False),
    ]


def run(monkeypatch, resolved, candidates=None, max_sources=10):
    pipeline = ResearchPipeline()
    for provider in (pipeline.openalex, pipeline.semantic_scholar, pipeline.crossref):
        monkeypatch.setattr(provider, "normalize_schema", lambda source: source.copy())
        monkeypatch.setattr(provider, "fetch_raw_sources", lambda *args: [])
    data = candidates if candidates is not None else pool()
    monkeypatch.setattr(pipeline.openalex, "fetch_raw_sources", lambda *args: data)
    return pipeline._discover_sources_real("Reasoning evaluation", max_sources, None, resolved)


def test_with_no_full_match_the_closest_matches_are_returned_and_labelled(monkeypatch):
    result = run(monkeypatch, resolve(SIX, None, paid=True))

    included = {s["uid"]: s for s in result["included"]}
    assert set(included) == {"a", "b", "c"}, "good sources are no longer thrown away because none is a review"
    assert included["a"]["criteria_result"] == "Met 4 of 5"
    assert LABEL_REVIEW in included["a"]["criteria_not_met"] and "not a review" in included["a"]["criteria_not_met"]
    assert included["c"]["criteria_result"] == "Met 3 of 5" and "not open access" in included["c"]["criteria_not_met"]
    assert "Criteria: Met 4 of 5; not met:" in included["a"]["inclusion_reason"]

    audit = result["audit"]
    assert audit["candidates_meeting_criteria"] == 0 and audit["closest_matches_included"] == 3
    assert audit["criteria"]["strict"] is False


def test_a_source_that_misses_too_many_is_excluded_and_says_how_many_it_met(monkeypatch):
    result = run(monkeypatch, resolve(SIX, None, paid=True))
    far = next(s for s in result["excluded"] if s["uid"] == "far")
    assert "Met only 1 of 5 selected criteria (at least 3 needed)" in far["exclusion_reason"]
    assert far["criteria_result"] == "Met 1 of 5" and "published in 2005" in far["criteria_not_met"]


def test_strict_mode_returns_nothing_here_and_says_why(monkeypatch):
    result = run(monkeypatch, resolve(SIX, None, paid=True, strict=True))
    assert result["included"] == []
    assert all("strict matching" in s["exclusion_reason"] for s in result["excluded"] if "criteria_result" in s)
    assert result["audit"]["criteria"]["strict"] is True


def test_full_matches_always_rank_above_closest_matches(monkeypatch):
    candidates = pool()
    candidates[2].update(citation_count=5000)   # a huge score, but it misses a criterion
    candidates.append(src(uid="rev", title="Reasoning evaluation systematic review", doi="10.1/rev", citation_count=11,
                          work_type="review"))
    result = run(monkeypatch, resolve(SIX, None, paid=True), candidates)
    assert result["included"][0]["uid"] == "rev" and result["included"][0]["criteria_result"] == "Met all 5"
    assert result["audit"]["candidates_meeting_criteria"] == 1


def test_closest_matches_only_fill_the_places_full_matches_leave(monkeypatch):
    candidates = pool() + [src(uid="rev", title="Reasoning evaluation systematic review", doi="10.1/rev", work_type="review")]
    result = run(monkeypatch, resolve(SIX, None, paid=True), candidates, max_sources=1)
    assert [s["uid"] for s in result["included"]] == ["rev"]
    assert result["audit"]["closest_matches_included"] == 0
    overflow = next(s for s in result["excluded"] if s["uid"] == "a")
    assert "Lower topical match" in overflow["exclusion_reason"] and overflow["criteria_result"] == "Met 4 of 5", \
        "even a source cut by rank shows how it fared against the criteria"


def test_the_type_selector_still_excludes_other_types_outright(monkeypatch):
    candidates = pool() + [src(uid="pre", title="Reasoning evaluation preprint", doi="10.1/pre", work_type="preprint")]
    result = run(monkeypatch, resolve([LABEL_CITED, LABEL_RECENT], "journal", paid=True), candidates)
    assert "pre" not in {s["uid"] for s in result["included"]}
    pre = next(s for s in result["excluded"] if s["uid"] == "pre")
    assert pre["exclusion_reason"].startswith("Did not meet a required criterion")


def test_the_defaults_are_adaptive_too(monkeypatch):
    candidates = [src(uid="old", title="Reasoning evaluation classic", doi="10.1/o", year=2015, citation_count=500),
                  src(uid="new", title="Reasoning evaluation new", doi="10.1/n", year=2026, citation_count=0)]
    result = run(monkeypatch, resolve([], None, paid=False), candidates)
    assert {s["uid"] for s in result["included"]} == {"old", "new"}, "a classic and a brand-new paper each meet 2 of the 3 defaults"


# --- the workbook and the response ----------------------------------------------------------------------------------------------------------

def test_the_workbook_has_the_criteria_columns_and_the_matching_mode(monkeypatch, tmp_path):
    pipeline = ResearchPipeline()
    pipeline.insight_engine.registry = type("R", (), {"generate_with_failover": lambda *a, **k: ("Summary.", "fake")})()
    for provider in (pipeline.openalex, pipeline.semantic_scholar, pipeline.crossref):
        monkeypatch.setattr(provider, "normalize_schema", lambda source: source.copy())
        monkeypatch.setattr(provider, "fetch_raw_sources", lambda *args: [])
    monkeypatch.setattr(pipeline.openalex, "fetch_raw_sources", lambda *args: pool())

    class Generator:
        def generate_comprehensive_dossier(self, **kwargs):
            dossier = ResearchDossier(kwargs["query"])
            dossier.included_sources = kwargs["included_sources"]
            dossier.themes = ["A theme"]
            return dossier

    monkeypatch.setattr("app.research.research_pipeline.DossierGenerator", Generator)
    result = pipeline.run_research("Reasoning evaluation", max_sources=10, output_directory=str(tmp_path),
                                   criteria=resolve(SIX, None, paid=True))
    book = load_workbook(tmp_path / result.report_data["discovery_report_name"], read_only=True)

    for sheet in ("Included Evidence", "Excluded Candidates"):
        headers = [str(h) for h in list(book[sheet].values)[0]]
        assert "Criteria Result" in headers and "Criteria Not Met" in headers, sheet
    summary = {str(k): v for k, v in list(book["Executive Summary"].values)[1:]}
    assert summary["Matching mode"] == "Closest matches fill the gap" and summary["Closest matches included"] == 3
    config = {str(k): v for k, v in list(book["Run Audit Configuration"].values)[1:]}
    assert config["Matching Mode"].startswith("Closest matches fill the gap")
    assert result.report_data["criteria"]["closest_matches_included"] == 3 and result.report_data["criteria"]["strict"] is False


def test_strict_is_stated_in_the_workbook(monkeypatch, tmp_path):
    pipeline = ResearchPipeline()
    pipeline.insight_engine.registry = type("R", (), {"generate_with_failover": lambda *a, **k: ("Summary.", "fake")})()
    monkeypatch.setattr(pipeline, "_discover_sources_real", lambda *args, **kwargs: {
        "included": [], "excluded": [], "audit": {"active_criteria": ["x"]}})
    result = pipeline.run_research("Reasoning evaluation", output_directory=str(tmp_path),
                                   criteria=resolve([LABEL_OPEN], None, True, strict=True))
    book = load_workbook(tmp_path / result.report_data["discovery_report_name"], read_only=True)
    assert {str(k): v for k, v in list(book["Executive Summary"].values)[1:]}["Matching mode"] == "Strict: every criterion required"


def test_the_free_preview_rows_keep_their_criteria_labels():
    row = exposure.preview_view({"included": [{"uid": "u", "title": "T", "criteria_result": "Met 4 of 5",
                                               "criteria_not_met": "Open access (not open access)", "abstract": "SECRET"}],
                                 "excluded": []})["included"][0]
    assert row["criteria_result"] == "Met 4 of 5" and "abstract" not in row


# --- the endpoint ---------------------------------------------------------------------------------------------------------------------------------

@pytest.fixture
def world(monkeypatch, tmp_path, fake_provider, synthesized_draft_json):
    return World(monkeypatch, tmp_path, fake_provider, synthesized_draft_json)


class Recording:
    def __init__(self):
        self.criteria = None

    def run_research(self, **kwargs):
        self.criteria = kwargs["criteria"]
        Path(kwargs["output_directory"], "research_audit_x_20261001_120000.xlsx").write_bytes(b"workbook")
        return {"status": "success", "included": sources(2), "excluded": [], "criteria": kwargs["criteria"].summary(),
                "discovery_report_name": "research_audit_x_20261001_120000.xlsx"}


@pytest.mark.parametrize("flag, expected", [(False, False), (True, True)])
def test_the_strict_switch_reaches_the_scan(world, monkeypatch, flag, expected):
    pipeline = Recording()
    monkeypatch.setattr(cloud_app, "pipeline", pipeline)
    response = asyncio.run(cloud_app.execute_cloud_scan(
        cloud_app.ScanRequest(topic="AI evidence", selected_inclusion_reasons=[LABEL_OPEN], strict_criteria=flag),
        authorization="Bearer t"))
    assert pipeline.criteria.strict is expected and response["criteria"]["strict"] is expected


def test_the_switch_is_off_by_default_and_not_counted_as_a_criterion(world, monkeypatch):
    pipeline = Recording()
    monkeypatch.setattr(cloud_app, "pipeline", pipeline)
    asyncio.run(cloud_app.execute_cloud_scan(
        cloud_app.ScanRequest(topic="AI evidence", selected_inclusion_reasons=[LABEL_OPEN, LABEL_CITED, LABEL_RECENT]),
        authorization="Bearer t"))
    assert pipeline.criteria.strict is False and len(pipeline.criteria.applied) == 3 and pipeline.criteria.limited is False


# --- the page ----------------------------------------------------------------------------------------------------------------------------------

import re

ROOT = Path(__file__).resolve().parent.parent
HTML = (ROOT / "website" / "index.html").read_text(encoding="utf-8")
JS = (ROOT / "website" / "app.js").read_text(encoding="utf-8")
CSS = (ROOT / "website" / "styles.css").read_text(encoding="utf-8")


def test_the_strict_switch_is_outside_the_criteria_so_it_never_counts_toward_the_limit():
    fieldset = HTML[HTML.index('<fieldset class="criteria-fieldset">'):HTML.index("</fieldset>")]
    assert 'id="strictCriteria"' not in fieldset, "inside the fieldset it would be sent as a criterion and counted"
    assert 'id="strictCriteria"' in HTML and "Require every selected criterion" in HTML


def test_the_switch_is_off_by_default_and_sent_with_the_scan():
    assert '<input type="checkbox" id="strictCriteria">' in HTML and "checked" not in HTML.split('id="strictCriteria"')[1][:20]
    assert "strict_criteria: byId('strictCriteria').checked" in JS


def test_every_list_says_what_it_is_based_on():
    assert 'id="criteriaUsed"' in HTML and 'id="includedBasis"' in HTML and 'id="excludedBasis"' in HTML
    body = JS[JS.index("function renderCriteriaUsed(data)"):]
    body = body[:body.index("\n    }\n")]
    assert "info.applied.join(' · ')" in body and "Strict matching: every criterion is required." in body
    assert "Closest matches fill the gap" in body and "innerHTML" not in body
    assert "renderCriteriaUsed(data);" in JS[JS.index("function renderResult("):][:200]


def test_included_sources_show_what_they_met_and_every_excluded_source_shows_why():
    body = JS[JS.index("function makeSourceCard(source, included)"):]
    body = body[:body.index("\n    }\n")]
    assert "source.criteria_result" in body and "criteria-chip partial" in body and "criteria-chip full" in body
    assert "if (!included)" in body and "Why excluded:" in body and "source-why" in body
    assert "innerHTML" not in body


def test_the_scan_message_reports_closest_matches():
    body = JS[JS.index("function criteriaNote(data)"):]
    body = body[:body.index("\n    }\n")]
    assert "closest_matches_included" in body and "closest" in body and 'switch off "Require every selected criterion"' in body


@pytest.mark.parametrize("name", ["criteria-strict", "criteria-used", "criteria-chip", "source-why"])
def test_the_new_classes_are_styled(name):
    assert f".{name}" in CSS
