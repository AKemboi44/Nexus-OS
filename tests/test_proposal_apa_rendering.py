"""APA formatting, justification and naming of rendered proposals (opens the real DOCX)."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from app.reports.apa import in_text_citation
from app.reports.proposal_service import document_filename, generate_proposal_docx

SOURCES = [
    {
        "title": "Fairness in automated resume screening",
        "authors": ["Zimmerman, Paula", "Adeyemi, Tunde", "Okafor, Ngozi"],
        "year": 2024,
        "venue": "Journal of Applied Machine Learning",
        "abstract": "Scalability remains a limitation of screening models. A novel approach is proposed.",
        "doi": "10.1000/jaml.2024.1",
    },
    {
        "title": "Evaluating language model agents",
        "authors": ["Baker, Chris", "Lee, Kim"],
        "year": 2023,
        "venue": "Computational Linguistics Review",
        "abstract": "Interpretability challenges limit adoption. Empirical validation shows promise.",
        "doi": "10.1000/clr.2023.2",
    },
    {
        "title": "Ethics of algorithmic hiring",
        "authors": ["Mensah, Ama"],
        "year": 2022,
        "venue": "Ethics and Information Technology",
        "abstract": "This systematic review examines hiring ethics and opportunities for practical deployment.",
        "doi": "10.1000/eit.2022.3",
    },
]


def _open(topic="Ethical AI Evaluation Frameworks"):
    document = generate_proposal_docx(topic, "scholarly", SOURCES)
    return document, Document(io.BytesIO(document.document_bytes))


def _split_body_and_references(doc):
    paragraphs = doc.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p.style.name == "Heading 1" and p.text == "References")
    return paragraphs[:index], paragraphs[index + 1:]


def test_body_paragraphs_are_justified_and_free_of_markdown_or_brackets():
    _, doc = _open()
    body, _ = _split_body_and_references(doc)
    prose = [p for p in body if p.style.name == "Normal" and p.text.strip()][1:]  # skip title
    assert len(prose) > 8
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY for p in prose)
    assert all(not p.paragraph_format.first_line_indent for p in prose), "paragraphs start flush left"
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "[" not in full_text and "]" not in full_text
    assert "*" not in full_text


def test_in_text_citations_follow_apa_author_date_style():
    _, doc = _open()
    body_text = " ".join(p.text for p in _split_body_and_references(doc)[0])
    assert re.search(r"\((?:[A-Z][a-z]+(?: & [A-Z][a-z]+| et al\.)?), \d{4}(?:; [^)]+)?\)", body_text)
    assert "Zimmerman et al., 2024" in body_text or "Baker & Lee, 2023" in body_text


def test_reference_list_is_alphabetical_with_hanging_indent_and_italics():
    _, doc = _open()
    _, references = _split_body_and_references(doc)
    entries = [p for p in references if p.text.strip()]
    assert entries
    surnames = [entry.text.split(",")[0] for entry in entries]
    assert surnames == sorted(surnames, key=str.casefold)
    for entry in entries:
        assert entry.paragraph_format.first_line_indent == Inches(-0.5)
        assert entry.paragraph_format.left_indent == Inches(0.5)
        assert any(run.italic for run in entry.runs)


def test_in_text_citation_groups_and_sorts_authors():
    assert in_text_citation([SOURCES[1]]) == "(Baker & Lee, 2023)"
    assert in_text_citation([SOURCES[0]]) == "(Zimmerman et al., 2024)"
    assert in_text_citation([SOURCES[2], SOURCES[1]]) == "(Baker & Lee, 2023; Mensah, 2022)"


def test_document_is_named_after_the_topic():
    document, _ = _open("Ethical AI: Evaluation/Frameworks for Resume Screening?")
    assert document.document_name == "Ethical AI Evaluation Frameworks for Resume Screening.docx"
    assert document_filename("") == "proposal.docx"
    assert document_filename("x" * 200).endswith(".docx")
    assert len(document_filename("x" * 200)) <= 85
