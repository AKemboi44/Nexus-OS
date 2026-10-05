"""The workbook reads like a deliverable: readable authors, a real contribution sentence, honest headers, no empty columns."""

import pytest
from openpyxl import load_workbook

from app.research.providers.openalex_provider import reconstruct_abstract
from app.research.research_pipeline import ResearchPipeline
from app.reports.excel_formatting import format_research_workbook
from models.research_dossier import ResearchDossier

LONG = (
    "Background and context about a broad area of work that readers already know well and that sets the scene. "
    "We found that models trained on curated data outperformed larger models trained on web text across the "
    "reasoning benchmarks that we evaluated in this study, with the largest gains on multi-step problems. "
    "Limitations and future work are discussed at the end of the paper in some detail."
)


# --- the contribution sentence --------------------------------------------------------------------------------------------

def contribution(abstract, **extra):
    return ResearchPipeline._source_contribution({"abstract": abstract, "title": "T", **extra})


def test_the_contribution_is_the_sentence_that_states_a_finding_not_the_scene_setting():
    text = contribution(LONG)
    assert text.startswith("We found that models trained on curated data")
    assert "Background and context" not in text


def test_a_long_finding_is_cut_at_a_word_boundary_never_mid_word():
    sentence = "We found that " + " ".join(f"measurement{n}" for n in range(60)) + " were all consistent."
    text = contribution(sentence)
    assert len(text) <= ResearchPipeline.CONTRIBUTION_CHARS + 1
    assert text.endswith("…") and not text[:-1].endswith((" ", ","))
    assert text[:-1].split(" ")[-1] in sentence.split(), "the last word is a whole word from the abstract"


def test_a_short_enough_finding_is_kept_whole_without_an_ellipsis():
    text = contribution(LONG)
    assert len(text) <= ResearchPipeline.CONTRIBUTION_CHARS and not text.endswith("…") and text.endswith("problems.")


def test_a_long_abstract_is_not_just_copied():
    text = contribution(LONG)
    assert text != LONG and len(text) < len(LONG) and "Limitations and future work" not in text


def test_short_abstracts_keep_their_first_sentences_whole():
    assert contribution("An evaluation of token optimization techniques in compiler pipelines.") == \
        "An evaluation of token optimization techniques in compiler pipelines."
    assert contribution("A short claim. A second short sentence follows here.") == "A short claim. A second short sentence follows here."


def test_with_no_finding_cue_it_uses_the_opening_sentence():
    text = contribution("Rivers carry sediment to the coast over many years. The delta grows slowly in most seasons.")
    assert text.startswith("Rivers carry sediment")


@pytest.mark.parametrize("abstract", ["", "   ", "No description.", None])
def test_no_abstract_falls_back_to_the_metadata_line(abstract):
    text = ResearchPipeline._source_contribution({"abstract": abstract, "title": "A Title", "venue": "A Journal", "year": 2024})
    assert text.startswith("Metadata contribution: A Title")
    assert "No description" not in text


# --- the abstract label ------------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("label, expected", [
    ("Abstract As artificial intelligence continues", "As artificial intelligence continues"),
    ("Abstract: Background text here", "Background text here"),
    ("Abstract. The study examined", "The study examined"),
    ("Abstract algebra is the study of structures", "Abstract algebra is the study of structures"),
    ("The abstract of this paper", "The abstract of this paper"),
])
def test_a_leading_abstract_heading_is_removed_but_real_sentences_are_not(label, expected):
    index = {}
    for position, word in enumerate(label.split()):
        index.setdefault(word, []).append(position)
    assert reconstruct_abstract(index) == expected


# --- the workbook ------------------------------------------------------------------------------------------------------------------------

def source(**extra):
    base = {"uid": "u1", "title": "Reasoning evaluation", "authors": ["Tahir S. Pillay", "Adil I. Khan", "Sedef Yenice"],
            "venue": "Journal", "year": 2025, "doi": "10.1/x", "url": "https://doi.org/10.1/x", "abstract": LONG,
            "citation_count": 36, "is_peer_reviewed": True, "is_open_access": True, "work_type": "journal-article",
            "domain": "scholarly", "provider_source": "openalex", "keywords": [], "influential_citations": 0,
            "inclusion_reason": "Relevant: matched reasoning.", "source_contribution": "x"}
    base.update(extra)
    return base


def build(monkeypatch, tmp_path, included, excluded=None):
    pipeline = ResearchPipeline()
    monkeypatch.setattr(pipeline, "_discover_sources_real", lambda *args, **kwargs: {
        "included": [dict(s) for s in included], "excluded": [dict(s) for s in (excluded or [])],
        "audit": {"candidates_retrieved": 3, "unique_candidates_reviewed": 3, "active_criteria": ["x"]}})
    pipeline.insight_engine.registry = type("R", (), {"generate_with_failover": lambda *a, **k: ("Summary.", "fake")})()

    class Generator:
        def generate_comprehensive_dossier(self, **kwargs):
            dossier = ResearchDossier(kwargs["query"])
            dossier.included_sources = kwargs["included_sources"]
            dossier.themes = ["A theme"]
            return dossier

    monkeypatch.setattr("app.research.research_pipeline.DossierGenerator", Generator)
    result = pipeline.run_research("Reasoning evaluation", max_sources=5, output_directory=str(tmp_path))
    return load_workbook(tmp_path / result.report_data["discovery_report_name"], read_only=True)


def rows(book, sheet):
    values = list(book[sheet].values)
    return [str(h) for h in values[0]], values[1:]


def test_authors_are_readable_text_in_the_sheet(monkeypatch, tmp_path):
    headers, data = rows(build(monkeypatch, tmp_path, [source()]), "Included Evidence")
    cell = data[0][headers.index("Authors")]
    assert cell == "Tahir S. Pillay, Adil I. Khan, Sedef Yenice" and "[" not in cell and "'" not in cell


def test_keywords_are_joined_when_there_are_some(monkeypatch, tmp_path):
    headers, data = rows(build(monkeypatch, tmp_path, [source(keywords=["ai", "testing"], influential_citations=3)]), "Included Evidence")
    assert data[0][headers.index("Keywords")] == "ai; testing"
    assert "Influential Citations" in headers


def test_optional_columns_that_are_empty_on_every_row_are_dropped(monkeypatch, tmp_path):
    headers, _ = rows(build(monkeypatch, tmp_path, [source()], [source(uid="u2", exclusion_reason="Why")]), "Included Evidence")
    assert "Keywords" not in headers and "Influential Citations" not in headers
    assert "Citation Count" in headers and "Abstract" in headers, "real data columns are never dropped"


def test_headers_say_what_they_mean(monkeypatch, tmp_path):
    headers, _ = rows(build(monkeypatch, tmp_path, [source()]), "Included Evidence")
    assert "Published in journal" in headers and "Is Peer Reviewed" not in headers
    assert "Open access" in headers and "Work Type" in headers and "Source ID" in headers


def test_the_contribution_column_is_the_key_sentence_not_the_abstract(monkeypatch, tmp_path):
    # the pipeline computes contributions for sources that arrive without one
    headers, data = rows(build(monkeypatch, tmp_path, [source(source_contribution=None)]), "Included Evidence")
    cell = data[0][headers.index("Source Contribution")]
    assert cell.startswith("We found that") and cell != LONG and len(cell) <= ResearchPipeline.CONTRIBUTION_CHARS + 1


def test_the_header_labels_apply_to_any_workbook(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "h.xlsx"
    book = Workbook()
    book.active.append(["is_peer_reviewed", "is_open_access", "citation_count"])
    book.save(path)
    format_research_workbook(path)
    assert [cell.value for cell in load_workbook(path).active[1]] == ["Published in journal", "Open access", "Citation Count"]
