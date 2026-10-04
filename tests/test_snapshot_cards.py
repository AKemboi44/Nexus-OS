"""The greyed Excel and Word preview cards: free users see what the files contain, the rest stays locked."""

import io
import re
from pathlib import Path

import pytest
from docx import Document

from app.analytics import schema
from app.reports.snapshot import snapshot_from_docx

ROOT = Path(__file__).resolve().parent.parent
JS = (ROOT / "website" / "app.js").read_text(encoding="utf-8")
HTML = (ROOT / "website" / "index.html").read_text(encoding="utf-8")
CSS = (ROOT / "website" / "styles.css").read_text(encoding="utf-8")
EXT_JS = (ROOT / "extension" / "popup.js").read_text(encoding="utf-8")
PIPELINE = (ROOT / "app" / "research" / "research_pipeline.py").read_text(encoding="utf-8")

START = JS.index("// --- free-tier preview cards")
CARDS = JS[START:JS.index("function renderUpsells(data)")]


def function_body(source, signature):
    body = source[source.index(signature):]
    return body[:body.index("\n    }\n")]


def test_the_excel_card_lists_exactly_the_sheets_the_workbook_really_has():
    real = set(re.findall(r"sheet_name=['\"]([^'\"]+)['\"]", PIPELINE))
    listed = set(re.findall(r"'([^']+)'", CARDS[CARDS.index("const EXCEL_SHEETS = ["):CARDS.index("];")]))
    assert listed == real and len(real) == 7, "the card must not describe a file that is not what we deliver"


def test_cards_show_only_to_free_users_with_a_preview_and_hide_the_downloads_then():
    assert "return !paid && Boolean(data && data.preview)" in CARDS
    body = function_body(CARDS, "function renderDeliverables(data)")
    assert "byId('downloadExcel').hidden = preview" in body and "byId('proposalReport').hidden = preview" in body
    assert "container.hidden = !preview" in body and "container.replaceChildren();" in body
    for button in ("downloadExcel", "proposalReport"):
        assert f'id="{button}"' in HTML, "paid users keep the real download buttons"


def test_cards_are_rebuilt_whenever_the_scan_or_the_entitlement_changes():
    assert "renderDeliverables(data);" in function_body(JS, "function renderResult(")
    assert "if (activeResult) renderDeliverables(activeResult);" in function_body(JS, "async function refreshEntitlement()")


def test_after_buying_a_pack_the_same_scan_is_reloaded_in_full_without_scrolling_away():
    body = function_body(JS, "async function finishCheckout(orderId)")
    assert "activeResult.preview" in body and "loadRun(activeResult.research_run_id, {scroll: false})" in body
    assert "if (scroll) byId('scanResults').scrollIntoView" in JS


def test_server_text_is_only_ever_rendered_as_text():
    assert "innerHTML" not in CARDS and "insertAdjacentHTML" not in CARDS
    assert "appendText(page, 'p', snap.opening_paragraph" in CARDS


def test_locked_areas_are_empty_placeholders_never_hidden_text():
    body = function_body(CARDS, "function ghostLines(count)")
    assert "textContent" not in body and "innerText" not in body and "aria-hidden" in body
    ghost_rows = CARDS[CARDS.index("row.className = 'sheet-row sheet-ghost'"):][:520]
    assert "textContent" not in ghost_rows and "aria-hidden" in ghost_rows
    assert ".sheet-ghost { filter: blur" in CSS and ".ghost-lines { display: grid" in CSS and "filter: blur" in CSS.split(".ghost-lines {")[1].split("}")[0]


def test_only_the_first_rows_the_server_sent_are_drawn_as_real_text():
    body = function_body(CARDS, "function excelCard(data)")
    assert "(data.included || []).slice(0, 3)" in body and "data.locked?.included_hidden" in body
    assert "data.excluded" not in body, "filtered candidates are described, not shown"


def test_the_word_preview_calls_the_server_once_per_click_with_the_saved_run():
    body = function_body(CARDS, "async function previewProposal()")
    assert "wordPreview.status === 'loading'" in body, "a second click while loading is ignored"
    assert "research_run_id: runId" in body and "preview: true" in body and "report_type: 'proposal'" in body
    assert "included_sources" not in body, "the server builds the report from the saved run"
    assert "detail.error_code === 'upgrade_required'" in body and "status: 'used'" in body


def test_every_word_state_has_a_next_step():
    body = function_body(CARDS, "function renderWordBody(body, data)")
    for state in ("'ready'", "'used'", "'error'", "'loading'"):
        assert state in body
    assert body.count("unlockButton(") >= 2 and "Try again" in body and "Preview my proposal (free)" in body
    assert "Writing your preview…" in body


def test_the_ui_reads_only_fields_the_server_snapshot_really_returns():
    document = Document()
    document.add_paragraph("A Title")
    document.add_heading("Research Overview", level=1)
    document.add_paragraph("word " * 30)
    document.add_heading("References", level=1)
    document.add_paragraph("Ref one")
    buffer = io.BytesIO()
    document.save(buffer)
    snapshot = snapshot_from_docx(buffer.getvalue())

    used = set(re.findall(r"\bsnap\.([a-z_]+)", CARDS)) | set(re.findall(r"\bsnapshot\.([a-z_]+)", EXT_JS))
    assert used and used <= set(snapshot), used - set(snapshot)
    section_keys = set(re.findall(r"\bsection\.([a-z_]+)", CARDS)) | set(re.findall(r"\bsection\.([a-z_]+)", EXT_JS))
    assert section_keys <= set(snapshot["sections"][0]), section_keys


def test_no_inline_styles_or_handlers_in_the_cards_markup():
    assert '<div class="snapshot-cards" id="snapshotCards" hidden></div>' in HTML
    assert "style=" not in CARDS and " onclick=" not in CARDS


def test_new_placements_and_the_preview_kind_are_registered():
    placements = schema.get("upgrade_cta_clicked").props["placement"].values
    for placement in re.findall(r"openUpgrade\('(snapshot_[a-z]+)'\)", CARDS):
        assert placement in placements, placement
    assert {"snapshot_excel", "snapshot_word", "extension_snapshot"} <= set(placements)
    assert schema.validate("deliverable_clicked", {"kind": "proposal_preview"}, schema.WEB) == []
    assert "track('deliverable_clicked', {kind: 'proposal_preview'})" in CARDS
    assert "openUpgradeFlow('extension_snapshot')" in EXT_JS


def test_the_card_layout_stacks_on_phones_and_cannot_widen_the_page():
    assert ".snapshot-cards { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr))" in CSS
    assert "@media (max-width: 880px) { .snapshot-cards { grid-template-columns: 1fr; } }" in CSS
    assert ".snapshot-outline { display: grid; grid-template-columns: minmax(0, 1fr)" in CSS
    assert ".snapshot-outline li { min-width: 0" in CSS


# --- the extension --------------------------------------------------------------------------------------------------------

def test_the_extension_asks_free_users_for_a_preview_and_paid_users_for_the_document():
    assert "preview: !isPaidUser && !isLocalTestingInstall && reportType === 'proposal'" in EXT_JS
    assert "...(request.preview ? {preview: true} : {})" in EXT_JS


def test_a_snapshot_is_handled_before_the_queued_job_check_that_would_misread_it():
    handler = EXT_JS[EXT_JS.index("async function handleBackgroundResponse(response)"):]
    assert handler.index("data.status === 'snapshot' && data.snapshot") < handler.index("const isQueuedReport")


def test_the_extension_renders_the_snapshot_as_text_with_an_unlock_button():
    body = EXT_JS[EXT_JS.index("function renderProposalSnapshot(snapshot)"):]
    body = body[:body.index("\n    }\n")]
    assert "innerHTML" not in body and "textContent" in body
    assert "openUpgradeFlow('extension_snapshot')" in body and "Unlock the full proposal" in body
