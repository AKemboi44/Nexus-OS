"""Model synthesis path: one call, strict gate, always falls back to a downloadable template."""

import io
import json

from docx import Document

from app.reports.proposal_service import AI_SYNTHESIZED, TEMPLATE_FALLBACK, generate_proposal_docx
from tests.test_proposal_apa_rendering import SOURCES


def _generate(provider):
    document = generate_proposal_docx("Ethical AI Evaluation Frameworks", "scholarly", SOURCES, provider=provider)
    assert Document(io.BytesIO(document.document_bytes)).paragraphs, "output must always be a real DOCX"
    return document


def test_valid_synthesis_is_used_with_a_single_provider_call(fake_provider, synthesized_draft_json):
    provider = fake_provider(content=synthesized_draft_json())
    document = _generate(provider)

    assert document.generation_mode == AI_SYNTHESIZED
    assert document.synthesis_note is None
    assert document.word_count >= 1000
    assert provider.calls == 1

    text = " ".join(p.text for p in Document(io.BytesIO(document.document_bytes)).paragraphs)
    assert "Zimmerman et al., 2024" in text and "Baker & Lee, 2023" in text


def _with_extra_overview_paragraphs(draft_json, count):
    draft = json.loads(draft_json)
    draft["overview"].extend(draft["overview"][0] for _ in range(count))
    return json.dumps(draft)


def test_overlong_but_valid_draft_is_kept(fake_provider, synthesized_draft_json):
    # The model returned 2,802 words in production; that must not discard a good synthesis.
    provider = fake_provider(content=_with_extra_overview_paragraphs(synthesized_draft_json(), 8))
    document = _generate(provider)

    assert document.generation_mode == AI_SYNTHESIZED
    assert 2000 < document.word_count < 3500


def test_runaway_length_still_falls_back(fake_provider, synthesized_draft_json):
    provider = fake_provider(content=_with_extra_overview_paragraphs(synthesized_draft_json(), 20))
    document = _generate(provider)

    assert document.generation_mode == TEMPLATE_FALLBACK
    assert document.synthesis_failure == "quality_gate"


def test_fenced_json_from_the_model_is_accepted(fake_provider, synthesized_draft_json):
    provider = fake_provider(content=f"Here is the proposal:\n```json\n{synthesized_draft_json()}\n```")
    assert _generate(provider).generation_mode == AI_SYNTHESIZED


def test_each_fallback_reports_its_specific_reason(fake_provider, synthesized_draft_json):
    class Unconfigured:
        name = "unconfigured"

        def is_configured(self):
            return False

    uncited = json.loads(synthesized_draft_json())
    uncited["research_gap"][0]["evidence_ids"] = []
    cases = {
        "not_configured": Unconfigured(),
        "provider_error": fake_provider(error=RuntimeError("401 invalid x-api-key")),
        "invalid_output": fake_provider(content="Sorry, I cannot do that."),
        "quality_gate": fake_provider(content=json.dumps(uncited)),
    }
    for reason, provider in cases.items():
        document = _generate(provider)
        assert document.synthesis_failure == reason
        assert document.synthesis_note


def test_draft_that_copies_source_wording_is_rejected(fake_provider, synthesized_draft_json):
    copied = json.loads(synthesized_draft_json())
    copied["overview"][0]["text"] += " " + SOURCES[0]["abstract"]
    provider = fake_provider(content=json.dumps(copied))
    document = _generate(provider)

    assert document.generation_mode == TEMPLATE_FALLBACK
    assert document.synthesis_note
    assert provider.calls == 1, "no retry: the budget is one call"


def test_provider_error_falls_back_without_retrying(fake_provider):
    provider = fake_provider(error=RuntimeError("529 overloaded"))
    assert _generate(provider).generation_mode == TEMPLATE_FALLBACK
    assert provider.calls == 1


def test_unparsable_output_falls_back(fake_provider):
    assert _generate(fake_provider(content="Sorry, I cannot do that.")).generation_mode == TEMPLATE_FALLBACK


def test_unknown_source_id_falls_back(fake_provider, synthesized_draft_json):
    draft = json.loads(synthesized_draft_json())
    draft["overview"][0]["evidence_ids"] = ["S1", "S9"]
    assert _generate(fake_provider(content=json.dumps(draft))).generation_mode == TEMPLATE_FALLBACK


def test_evidence_paragraph_without_citation_falls_back(fake_provider, synthesized_draft_json):
    draft = json.loads(synthesized_draft_json())
    draft["research_gap"][0]["evidence_ids"] = []
    assert _generate(fake_provider(content=json.dumps(draft))).generation_mode == TEMPLATE_FALLBACK


def test_unconfigured_provider_falls_back_without_a_call():
    class Unconfigured:
        name = "unconfigured"
        calls = 0

        def is_configured(self):
            return False

    document = _generate(Unconfigured())
    assert document.generation_mode == TEMPLATE_FALLBACK
    assert document.synthesis_note
