from openpyxl import load_workbook
from openpyxl import Workbook

from app.research.research_pipeline import ResearchPipeline
from app.reports.excel_formatting import format_research_workbook
from models.research_dossier import ResearchDossier
from models.research_insight import ResearchInsight


def test_research_workbook_exports_research_areas_and_topic_opportunities(
    monkeypatch, tmp_path
):
    source = {
        "uid": "source-1",
        "title": "Research quality across community studies",
        "abstract": (
            "This study compares quality measures across community research settings "
            "and discusses limitations in recruitment and reporting."
        ),
        "authors": ["Researcher, A."],
        "year": 2024,
        "url": "https://example.org/study",
        "doi": "10.1000/example",
    }
    pipeline = ResearchPipeline()
    monkeypatch.setattr(
        pipeline,
        "_discover_sources_real",
        lambda *args: {"included": [source.copy()], "excluded": []},
    )
    monkeypatch.setattr(
        pipeline.insight_engine,
        "generate_insight",
        lambda *args: ResearchInsight("topic", "A concise synthesis.", []),
    )

    class Generator:
        def generate_comprehensive_dossier(self, **kwargs):
            dossier = ResearchDossier(kwargs["query"])
            dossier.included_sources = kwargs["included_sources"]
            dossier.abstract = "A grounded synthesis of research quality."
            dossier.themes = ["Quality measurement differs across settings."]
            dossier.research_areas = [
                "Compare recruitment methods across the community studies in the evidence base."
            ]
            dossier.opportunity_areas = [
                "Evaluate a shared reporting toolkit for community research teams."
            ]
            return dossier

    monkeypatch.setattr("app.research.research_pipeline.DossierGenerator", Generator)
    result = pipeline.run_research("research quality", output_directory=str(tmp_path))
    report_path = tmp_path / result.report_data["discovery_report_name"]

    workbook = load_workbook(report_path, read_only=False)
    assert "Research Areas" in workbook.sheetnames
    assert "Research Opportunities" in workbook.sheetnames
    research_sheet = workbook["Research Areas"]
    opportunities_sheet = workbook["Research Opportunities"]
    assert research_sheet["A2"].value.startswith("Compare recruitment methods")
    assert opportunities_sheet["A2"].value.startswith("Evaluate a shared reporting toolkit")
    assert research_sheet.freeze_panes == "A2"
    assert research_sheet["A1"].font.bold is True
    for worksheet in workbook.worksheets:
        assert all(
            str(cell.value)[0].isupper()
            for cell in worksheet[1]
            if cell.value not in (None, "")
        )
    workbook.close()


def test_workbook_capitalizes_headers_and_preserves_common_acronyms(tmp_path):
    path = tmp_path / "headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["research_topic", "doi", "source URL", "apa_format_citation"])
    workbook.save(path)

    format_research_workbook(path)

    formatted = load_workbook(path, read_only=True)
    assert list(formatted.active.values)[0] == (
        "Research Topic",
        "DOI",
        "Source URL",
        "APA Format Citation",
    )
    formatted.close()
