from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


_HEADER_ACRONYMS = {
    "ai", "api", "apa", "doi", "id", "isbn", "issn", "orcid", "pdf", "url",
}


def _title_case_header(value):
    words = str(value or "").replace("_", " ").split()
    return " ".join(
        word.upper() if word.casefold() in _HEADER_ACRONYMS
        else word[:1].upper() + word[1:].lower()
        for word in words
    )


def format_research_workbook(path):
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        if sheet.max_row and sheet.max_column:
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.value = _title_case_header(cell.value)
                cell.fill = PatternFill("solid", fgColor="233B63")
                cell.font = Font(name="Aptos", color="FFFFFF", bold=True)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 30

            sample_end = min(sheet.max_row, 201)
            for column in sheet.iter_cols(min_row=1, max_row=sample_end):
                longest = max(
                    (len(str(cell.value or "")) for cell in column),
                    default=0,
                )
                sheet.column_dimensions[column[0].column_letter].width = min(
                    max(longest + 2, 14),
                    58,
                )
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)
