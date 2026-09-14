# app/agents/scribe_agent.py - Part A
import os
import sys
import re
import time
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.synthesis.citation_engine import CitationEngine


class ScribeResearchAgent:
    """
    Advanced Academic Drafting Agent (Post-MVP Engine)
    Generates structured Word documents with cohesive grouped citation patterns.
    """

    def generate_apa_dossier_report(self, topic: str, included_sources: list,
                                    dossier=None, domain: str = "scholarly") -> str:
        return self._generate_report(
            topic, included_sources, dossier=dossier, domain=domain,
            report_type="proposal"
        )

    def generate_full_research_starter_report(self, topic: str, included_sources: list,
                                               dossier=None, domain: str = "scholarly") -> str:
        return self._generate_report(
            topic, included_sources, dossier=dossier, domain=domain,
            report_type="full_starter"
        )

    def _generate_report(self, topic: str, included_sources: list,
                         dossier=None, domain: str = "scholarly",
                         report_type: str = "proposal") -> str:
        doc = Document()

        # Configure standard academic 1-inch margins on all sides
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Enforce standard APA Times New Roman 12pt typography
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # 1. Inject Standard APA Cover Title Page Elements
        title_p = doc.add_paragraph()
        title_p.alignment = 1  # Center Alignment Block
        title_p.paragraph_format.space_before = Pt(120)

        report_title = (
            "Comprehensive Pre-Research Proposal Report"
            if report_type == "proposal"
            else "Full Research Starter Report"
        )
        title_run = title_p.add_run(f"{report_title}:\n{topic.title()}\n")
        title_run.bold = True

        author_p = doc.add_paragraph()
        author_p.alignment = 1
        author_p.paragraph_format.space_before = Pt(24)
        author_p.add_run(
            "Abraham Kiptoo Kemboi\n"
            "Nexus Research AI Collective Core\n"
            "Post-MVP Automation Node Suite"
        )

        doc.add_page_break()
        return self._write_synthesis_body(
            doc, topic, included_sources, dossier=dossier, domain=domain,
            report_type=report_type
        )

    # app/agents/scribe_agent.py - Part B
    def _write_synthesis_body(self, doc: Document, topic: str, included_sources: list,
                              dossier=None, domain: str = "scholarly",
                              report_type: str = "proposal") -> str:
        # 2. Main Title Header (Heading 1 Level)
        h1 = doc.add_paragraph()
        h1.alignment = 1
        h1_run = h1.add_run(
            "Literature Review and Research Starter Synthesis"
            if report_type == "full_starter"
            else "Literature Review and Pre-Research Proposal Synthesis"
        )
        h1_run.bold = True
        h1.paragraph_format.space_after = Pt(18)

        abstract = getattr(dossier, "abstract", "") if dossier else ""
        if not abstract:
            abstract = (
                f"This report synthesizes {len(included_sources)} selected {domain} source(s) "
                f"concerning {topic}."
            )
        self._add_heading(doc, "Abstract")
        self._add_body_paragraph(
            doc,
            f"This is an initial evidence-based research draft for researcher review. {abstract} "
            f"{self._citation_for_sources(included_sources)}"
        )

        themes = getattr(dossier, "themes", []) if dossier else []
        contradictions = getattr(dossier, "contradictions", []) if dossier else []
        gaps = getattr(dossier, "research_gaps", []) if dossier else []
        opportunities = getattr(dossier, "opportunity_areas", []) if dossier else []
        problems = getattr(dossier, "problems_to_solve", []) if dossier else []

        if themes:
            self._add_heading(doc, "Key Themes")
            for theme in themes:
                self._add_synthesis_paragraph(
                    doc, theme, included_sources,
                    f"Taken together, this theme frames the direction of the proposed study on {topic}."
                )

        if contradictions:
            self._add_heading(doc, "Contradictions and Boundary Conditions")
            for item in contradictions:
                self._add_synthesis_paragraph(
                    doc, item, included_sources,
                    "Nevertheless, these differences indicate that the finding should be interpreted within its specific context."
                )

        if gaps:
            self._add_heading(doc, "Research Gaps")
            for item in gaps:
                self._add_synthesis_paragraph(
                    doc, item, included_sources,
                    "This gap provides a direct justification for the proposed research."
                )

        if opportunities:
            self._add_heading(doc, "Opportunity Areas")
            for item in opportunities:
                self._add_synthesis_paragraph(
                    doc, item, included_sources,
                    "Moreover, this opportunity suggests a practical direction for subsequent investigation."
                )

        if problems:
            self._add_heading(doc, "Research Problems")
            for item in problems:
                self._add_synthesis_paragraph(
                    doc, item, included_sources,
                    "Overall, this research problem links the existing evidence to the study that should follow."
                )

        # 4. References Page Layout Module (Hanging Indent)
        doc.add_page_break()
        ref_h = doc.add_paragraph()
        ref_h.alignment = 1
        ref_h_run = ref_h.add_run("References")
        ref_h_run.bold = True
        ref_h.paragraph_format.space_after = Pt(12)

        for src in included_sources:
            ref_p = doc.add_paragraph()
            ref_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            ref_p.paragraph_format.line_spacing = 2.0
            ref_p.paragraph_format.space_after = Pt(0)
            ref_p.paragraph_format.left_indent = Inches(0.5)
            ref_p.paragraph_format.first_line_indent = Inches(-0.5)

            source_dict = src if isinstance(src, dict) else vars(src)
            ref_p.add_run(CitationEngine.generate_apa_7th(source_dict).replace("*", ""))

        filename_docx = (
            f"{'comprehensive_pre_research_proposal_report' if report_type == 'proposal' else 'full_research_starter_report'}_{self._clean_filename(topic)}_"
            f"{time.strftime('%Y%m%d_%H%M%S')}.docx"
        )
        target_docx_path = os.path.join(r"C:\Users\Abraham.Kemboi\PycharmProjects\Nexus-os", filename_docx)
        doc.save(target_docx_path)
        return target_docx_path

    @staticmethod
    def _add_heading(doc, text):
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(0)
        heading.add_run(text).bold = True

    @staticmethod
    def _add_body_paragraph(doc, text):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = Inches(0)
        clean_text = (
            str(text)
            .replace("**", "")
            .replace("*", "")
            .replace("—", ",")
            .replace("–", "-")
            .strip()
        )
        paragraph.add_run(clean_text)

    def _add_synthesis_paragraph(self, doc, evidence_text, sources, conclusion):
        citation = self._citation_for_sources(sources)
        text = (
            f"{self._topic_sentence(evidence_text)} "
            f"{citation} "
            f"Moreover, the combined evidence indicates that {self._interpretation(evidence_text)} "
            f"{conclusion}"
        )
        self._add_body_paragraph(doc, text)

    @staticmethod
    def _topic_sentence(text):
        sentence = str(text).strip()
        if not sentence:
            return "The available literature identifies a meaningful pattern in this area."
        return sentence[0].upper() + sentence[1:]

    @staticmethod
    def _interpretation(text):
        sentence = str(text).strip().rstrip(".")
        return f"this pattern should be understood in relation to {sentence.lower()}"

    @staticmethod
    def _source_value(source, key, default=""):
        if isinstance(source, dict):
            value = source.get(key, default)
        else:
            value = getattr(source, key, default)
        if isinstance(value, list):
            return ", ".join(str(item) for item in value)
        return value if value not in (None, "") else default

    def _citation_for_sources(self, sources):
        citations = []
        for source in sources:
            authors = self._source_value(source, "authors", "Unknown Author")
            year = self._source_value(source, "year", "n.d.")
            if isinstance(authors, list):
                author_text = str(authors[0]) if authors else "Unknown Author"
            else:
                author_text = str(authors)
            surname = author_text.replace(";", ",").split(",")[0].split()[-1]
            citations.append(f"{surname}, {year}")
        return f"({'; '.join(citations)})" if citations else ""

    def _clean_filename(self, query: str) -> str:
        return re.sub(r'[^a-zA-Z0-9]', '_', query.strip().lower())[:30]
